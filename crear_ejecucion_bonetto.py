from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# Configurar márgenes
section = doc.sections[0]
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin = Cm(3)
section.right_margin = Cm(2.5)

# Estilo base
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)

def add_paragraph(text='', bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, space_before=0, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
    return p

def add_mixed(parts, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6):
    """parts = lista de (texto, bold)"""
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    for text, bold in parts:
        run = p.add_run(text)
        run.bold = bold
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
    return p

# ── ENCABEZADO ──────────────────────────────────────────────────────────────
add_paragraph(
    'OBJETO: PROMUEVO INCIDENTE DE EJECUCIÓN DE SENTENCIA.',
    bold=True,
    align=WD_ALIGN_PARAGRAPH.CENTER,
    space_after=12
)

add_paragraph('SEÑOR JUEZ:', bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=10)

# ── PRESENTACIÓN ────────────────────────────────────────────────────────────
add_mixed([
    ('\t', False),
    ('BONETTO ALEXANDRA LEONELA', True),
    (', D.N.I. 43.149.095, con domicilio real en la calle Las Piedras 320 de '
     'la ciudad de Victoria, Provincia de Entre Ríos, por derecho propio y con '
     'el patrocinio letrado del Dr. Fernando Martín Juárez, Mat. 9811 T° I F° 266 CAER, '
     'constituyendo domicilio procesal en calle Mitre 500 de esta Ciudad, en autos ', False),
    ('"BONETTO ALEXANDRA LEONELA C/ BARANDELLI ANTONIO DAVID S/ ORDINARIO DAÑOS '
     'Y PERJUICIOS", Expediente N° 17416', True),
    (', a V.S. respetuosamente me presento y digo:', False),
], space_after=12)

# ── I. OBJETO ───────────────────────────────────────────────────────────────
add_paragraph('\t\tI.- OBJETO.', bold=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_mixed([
    ('\t', False),
    ('Que vengo a iniciar formal acción de ', False),
    ('INCIDENTE DE EJECUCIÓN DE SENTENCIA', True),
    (', en virtud de lo establecido por los artículos 485, 486 inc. 3°) y concordantes '
     'del C.P.C.C.E.R. (Ley N° 5315), contra ', False),
    ('BARANDELLI ANTONIO DAVID', True),
    (', D.N.I. 27.136.539, con domicilio en la calle Laprida e/ Monte Caseros y '
     'Constitución de la ciudad de Victoria, Provincia de Entre Ríos, por cobro de la suma de ', False),
    ('PESOS UN MILLÓN CUATROCIENTOS SETENTA Y CUATRO MIL NOVECIENTOS ($1.474.900)', True),
    (' —capital de condena determinado a valores actuales a la fecha de la sentencia del '
     '20 de abril de 2026—, con más la actualización por Índice de Precios al Consumidor '
     '(IPC) publicado por el INDEC computada desde la fecha de la presente sentencia hasta '
     'el efectivo pago, y la tasa de interés pura del ', False),
    ('TRES POR CIENTO (3%) ANUAL', True),
    (' calculada bajo el sistema de interés simple desde la notificación de la sentencia '
     'hasta el efectivo pago, más costas, o el monto que en más o en menos resulte de autos, '
     'correspondiente al crédito según Sentencia firme de fecha 20 de abril de 2026.', False),
], space_after=12)

# ── II. HECHOS ───────────────────────────────────────────────────────────────
add_paragraph('\t\tII.- HECHOS – SENTENCIA – MONTO RECLAMADO.', bold=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_mixed([
    ('\t', False),
    ('En los autos principales caratulados ', False),
    ('"BONETTO ALEXANDRA LEONELA C/ BARANDELLI ANTONIO DAVID S/ ORDINARIO DAÑOS Y PERJUICIOS", '
     'Expediente N° 17416', True),
    (', el Sr. Juez de Primera Instancia Dr. Luis Francisco Márquez Chada dictó Sentencia '
     'Definitiva con fecha ', False),
    ('20 de abril de 2026', True),
    (', mediante la cual resolvió:', False),
], space_after=8)

add_mixed([
    ('\t', False),
    ('1°) ', True),
    ('Hacer lugar a la demanda de daños y perjuicios promovida por BONETTO ALEXANDRA LEONELA '
     'contra BARANDELLI ANTONIO DAVID, condenando a este último a abonar a la actora, '
     'en el plazo de ', False),
    ('DIEZ (10) DÍAS', True),
    (' de notificada la presente, la suma de ', False),
    ('PESOS UN MILLÓN CUATROCIENTOS SETENTA Y CUATRO MIL NOVECIENTOS ($1.474.900)', True),
    (', compuesta por: a) Daño material (valor de reposición): $1.294.900; '
     'b) Privación de uso: $180.000.', False),
], space_after=8)

add_mixed([
    ('\t', False),
    ('2°) ', True),
    ('Declarar de oficio la inconstitucionalidad e inconvencionalidad de los arts. 7 y 10 '
     'de la Ley 23.928 y del art. 4 de la Ley 25.561, estableciendo que el capital de condena '
     'será actualizado mediante la aplicación del ', False),
    ('Índice de Precios al Consumidor (IPC) publicado por el INDEC', True),
    (', computándose la variación entre la fecha de la sentencia y el efectivo pago, con más '
     'una ', False),
    ('tasa de interés pura del 3% ANUAL', True),
    (' simple desde la notificación de la sentencia hasta el efectivo pago.', False),
], space_after=8)

add_mixed([
    ('\t', False),
    ('3°) ', True),
    ('Imponer las costas del proceso al demandado vencido (art. 65 del C.P.C.C.E.R.).', False),
], space_after=8)

add_mixed([
    ('\t', False),
    ('4°) ', True),
    ('Diferir la regulación de honorarios profesionales para la oportunidad prevista en el '
     'art. 160, inc. 8° del C.P.C.C.E.R., una vez practicada y aprobada la liquidación final.', False),
], space_after=10)

add_mixed([
    ('\t', False),
    ('A pesar del tiempo transcurrido desde la notificación de la sentencia y del vencimiento '
     'del plazo de diez (10) días fijado para el pago, el demandado ', False),
    ('BARANDELLI ANTONIO DAVID', True),
    (' —quien fuera declarado rebelde en los presentes autos—  no ha abonado hasta el '
     'presente suma alguna en concepto del crédito reconocido por la sentencia firme, '
     'por lo que ha incurrido en mora y corresponde en derecho la presente acción a fin '
     'de percibir las acreencias debidas a mi mandante.', False),
], space_after=10)

add_mixed([
    ('\t', False),
    ('Al precitado importe deberán adicionársele la actualización por IPC-INDEC y los '
     'intereses del 3% anual hasta el día del completo y efectivo pago, más las costas '
     'del presente incidente.', False),
], space_after=12)

# ── III. DERECHO ─────────────────────────────────────────────────────────────
add_paragraph('\t\tIII.- DERECHO.', bold=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_mixed([
    ('\t', False),
    ('Fundamos el derecho de esta parte en los artículos ', False),
    ('485, 486, 487, 488, 491 y concordantes del C.P.C.C.E.R. (Ley N° 5315)', True),
    (', artículos 1737, 1740, 1757, 1758 y 1769 del Código Civil y Comercial de la Nación, '
     'jurisprudencia y doctrina aplicables al caso.', False),
], space_after=12)

# ── IV. EMBARGO ───────────────────────────────────────────────────────────────
add_paragraph('\t\tIV.- EMBARGO EJECUTORIO.', bold=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_mixed([
    ('\t', False),
    ('De conformidad con el art. 488 y concordantes del C.P.C.C.E.R., se ordenará la '
     'traba de ', False),
    ('EMBARGO EJECUTORIO', True),
    (' hasta cubrir la suma reclamada, con más la que se presupueste provisoriamente '
     'para intereses y costas, sobre los bienes del demandado ', False),
    ('BARANDELLI ANTONIO DAVID', True),
    (', D.N.I. 27.136.539.', False),
], space_after=8)

add_mixed([
    ('\t', False),
    ('A tal efecto, se solicita:', False),
], space_after=6)

add_mixed([
    ('\t', False),
    ('a) ', True),
    ('Se libre oficio al Registro Nacional de la Propiedad del Automotor (Seccional '
     'correspondiente a la ciudad de Victoria, Provincia de Entre Ríos), a fin de '
     'que informe si el Sr. BARANDELLI ANTONIO DAVID, D.N.I. 27.136.539, figura '
     'como titular registral de rodados, y en su caso se trabe embargo ejecutorio '
     'sobre los mismos.', False),
], space_after=6)

add_mixed([
    ('\t', False),
    ('b) ', True),
    ('Se libre oficio al Registro de la Propiedad Inmueble de la Provincia de Entre Ríos, '
     'a fin de que informe si el demandado figura como titular de bienes inmuebles, '
     'trabándose en su caso el embargo ejecutorio correspondiente.', False),
], space_after=6)

add_mixed([
    ('\t', False),
    ('c) ', True),
    ('Subsidiariamente, en caso de no constatarse bienes suficientes, se libre oficio al '
     'Banco Central de la República Argentina (BCRA) a efectos de que informe la '
     'existencia de cuentas bancarias y/o depósitos a nombre del demandado, '
     'ordenándose en su caso la retención de los fondos necesarios.', False),
], space_after=6)

add_mixed([
    ('\t', False),
    ('d) ', True),
    ('Se libre ', False),
    ('INHIBICIÓN GENERAL DE BIENES', True),
    (' sobre el Sr. BARANDELLI ANTONIO DAVID, D.N.I. 27.136.539, con carácter '
     'subsidiario o complementario a las medidas precedentes, hasta cubrir el monto '
     'reclamado, sus intereses, actualización y costas.', False),
], space_after=10)

add_mixed([
    ('\t', False),
    ('En los despachos a librarse se dejará constancia de que el Dr. Fernando Martín '
     'Juárez se encuentra facultado para intervenir en el diligenciamiento de los '
     'mismos con las más amplias facultades de ley.', False),
], space_after=12)

# ── V. CITACION DE VENTA ─────────────────────────────────────────────────────
add_paragraph('\t\tV.- CITACIÓN DE VENTA.', bold=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_mixed([
    ('\t', False),
    ('Conforme al art. 491 y concordantes del C.P.C.C.E.R., se citará de venta al '
     'ejecutado ', False),
    ('BARANDELLI ANTONIO DAVID', True),
    (' por el plazo y bajo los apercibimientos de ley.', False),
], space_after=12)

# ── VI. DOCUMENTAL ────────────────────────────────────────────────────────────
add_paragraph('\t\tVI.- DOCUMENTAL.', bold=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_mixed([
    ('\t', False),
    ('De conformidad con lo normado por el art. 12 del Reglamento de Presentaciones '
     'Electrónicas, se indican las piezas necesarias de la causa principal para su '
     'exportación a este incidente:', False),
], space_after=8)

add_mixed([
    ('\t', False),
    ('1) ', True),
    ('Sentencia Definitiva de fecha 20 de abril de 2026, dictada por el Dr. Luis '
     'Francisco Márquez Chada, Juez de Primera Instancia, que hace lugar a la demanda '
     'y establece el crédito aquí ejecutado.', False),
], space_after=6)

add_mixed([
    ('\t', False),
    ('2) ', True),
    ('Constancias de notificación de la sentencia al demandado rebelde conforme '
     'art. 59 del C.P.C.C.E.R.', False),
], space_after=12)

# ── VII. EXCEPCION MEDIACION ──────────────────────────────────────────────────
add_paragraph('\t\tVII.- EXCEPCIÓN AL PROCESO DE MEDIACIÓN.', bold=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_mixed([
    ('\t', False),
    ('Conforme lo dispuesto por el art. 286 Bis in fine del C.P.C.C.E.R., optamos '
     'por excluir a este proceso de la mediación previa obligatoria señalada por '
     'dicha normativa, por tratarse de un incidente de ejecución de sentencia '
     'dictada en estos mismos autos.', False),
], space_after=12)

# ── VIII. PETITORIO ───────────────────────────────────────────────────────────
add_paragraph('\t\tVIII.- PETITORIO.', bold=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)

add_paragraph('\tPor lo expuesto, de V.S. SOLICITAMOS:', space_after=6)

peticiones = [
    ('1.- ', 'Me tenga por presentada, por denunciado domicilio real, manteniendo el domicilio '
     'procesal constituido en autos.'),
    ('2.- ', 'Tenga por iniciado Incidente de Ejecución de Sentencia contra el accionado '
     'BARANDELLI ANTONIO DAVID, D.N.I. 27.136.539, con el domicilio real denunciado, por '
     'cobro de la suma de PESOS UN MILLÓN CUATROCIENTOS SETENTA Y CUATRO MIL NOVECIENTOS '
     '($1.474.900) —capital de condena a valores actuales a la fecha de la sentencia—, '
     'con más la actualización por IPC-INDEC desde el 20/04/2026 y los intereses puros '
     'del 3% anual simple desde la notificación de la sentencia, más costas.'),
    ('3.- ', 'Tenga por acompañada en soporte papel copia para traslado de la presente acción.'),
    ('4.- ', 'Ordene el embargo ejecutorio interesado en el Capítulo IV del presente, '
     'librando los pertinentes oficios y el mandamiento de embargo solicitados.'),
    ('5.- ', 'Conforme lo interesado en el Capítulo V, cite de venta al ejecutado por '
     'el plazo y bajo los apercibimientos legales.'),
    ('6.- ', 'Atento a lo manifestado en el Capítulo VII, tenga por formulada la opción '
     'de excluir a esta acción del proceso de mediación previa obligatoria.'),
    ('7.- ', 'Oportunamente, mande llevar adelante la ejecución en todas sus partes, hasta '
     'que la actora se haga pago íntegro de sus acreencias, con más intereses, '
     'actualización y costas.'),
]

for num, texto in peticiones:
    add_mixed([
        ('\t', False),
        (num, True),
        (texto, False),
    ], space_after=6)

add_paragraph('', space_after=6)
add_paragraph('\tProveer de conformidad, SERÁ JUSTICIA.-', space_after=24)

# ── FIRMAS ────────────────────────────────────────────────────────────────────
add_paragraph('BONETTO ALEXANDRA LEONELA', bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
add_paragraph('D.N.I. 43.149.095', bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=16)

add_paragraph('Fernando Martín Juárez', bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
add_paragraph('Mat. 9811 T° I F° 266 CAER', bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=16)

add_mixed([
    ('Manifiesto que la patrocinada está en fiel conocimiento del contenido del presente '
     'escrito, conservando en mi poder una copia firmada del mismo, cumpliendo así con '
     'el Reglamento de Presentaciones Electrónicas.', False),
], space_after=0)

# ── GUARDAR ───────────────────────────────────────────────────────────────────
output = '/home/user/estudiojuarez/ejecucion_sentencia_bonetto_barandelli.docx'
doc.save(output)
print(f'Documento guardado en: {output}')
